"""
행정안전부 공공서비스(혜택) 데이터 수집 · 전처리 · DB 적재 파이프라인
=====================================================================
수정 사항:
  - [BUG FIX] upsert: if_exists="append" 단독 사용 → DELETE + INSERT 방식으로 교체
    (기존 코드는 실행할 때마다 동일 데이터가 무한 중복 적재되는 버그 존재)
"""

import re
import logging
from datetime import datetime

import requests
import pandas as pd
from sqlalchemy import create_engine, text
from apscheduler.schedulers.blocking import BlockingScheduler

# ──────────────────────────────────────────────
# 설정
# ──────────────────────────────────────────────
API_KEY  = "6c89345579052dd1b4441c2375af17384af98483ce10258fc3c5ac97ec82d72d"
ENDPOINT = "https://api.odcloud.kr/api/gov24/v3/serviceList"

DB_URL      = "sqlite:///gov_benefits.db"
TABLE_NAME  = "gov_benefits"
PK_COLUMN   = "서비스명"
SAVE_EXCEL  = True
SCHEDULE_HR = 2

DROP_COLUMNS = {"등록일시", "수정일시", "소관기관코드", "부서명", "서비스ID"}

# ──────────────────────────────────────────────
# 로깅
# ──────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
log = logging.getLogger(__name__)


# ──────────────────────────────────────────────
# 1. 수집
# ──────────────────────────────────────────────
def fetch_all() -> list[dict]:
    """API 페이지네이션으로 전체 데이터 수집"""
    all_rows, page, per_page = [], 1, 100
    log.info("데이터 수집 시작")

    while True:
        res = requests.get(
            ENDPOINT,
            params={"serviceKey": API_KEY, "page": page, "perPage": per_page},
            timeout=15,
        )
        res.raise_for_status()
        body  = res.json()
        rows  = body.get("data", [])
        total = body.get("totalCount", 0)

        if not rows:
            break

        all_rows.extend(rows)
        log.info("  페이지 %d 수집 완료 (%d / %d 건)", page, len(all_rows), total)

        if len(all_rows) >= total:
            break
        page += 1

    log.info("수집 완료: 총 %d 건", len(all_rows))
    return all_rows


# ──────────────────────────────────────────────
# 2. 전처리 (pandas)
# ──────────────────────────────────────────────
CONTROL_CHAR_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

def _clean_str(s: str) -> str:
    return CONTROL_CHAR_RE.sub("", s).strip()

def preprocess(raw: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame(raw)
    log.info("원본 컬럼 수: %d", len(df.columns))

    # 1) 불필요 컬럼 제거
    drop_existing = DROP_COLUMNS & set(df.columns)
    df = df.drop(columns=list(drop_existing))
    log.info("제거 컬럼: %s", drop_existing)

    # 2) 문자열 정제
    str_cols = df.select_dtypes(include="object").columns
    df[str_cols] = df[str_cols].apply(
        lambda col: col.map(lambda v: _clean_str(str(v)) if pd.notna(v) else v)
    )

    # 3) 날짜 컬럼 변환
    date_pattern = re.compile(r"일자|날짜|기간|시작|종료|일시", re.IGNORECASE)
    for col in df.columns:
        if date_pattern.search(col):
            df[col] = pd.to_datetime(df[col], errors="coerce")

    # 4) 결측값 정규화
    df = df.replace({"": pd.NA, "None": pd.NA, "null": pd.NA, "-": pd.NA})

    # 5) PK 기준 중복 제거 (가장 마지막 값 유지)
    if PK_COLUMN in df.columns:
        before = len(df)
        df = df.drop_duplicates(subset=[PK_COLUMN], keep="last")
        log.info("중복 제거: %d → %d 행", before, len(df))

    # 6) 컬럼명 소문자 + 공백→언더스코어 (DB 친화적)
    df.columns = [
        col.strip()
           .lower()
           .replace(" ", "_")
           .replace("(", "")
           .replace(")", "")
        for col in df.columns
    ]

    log.info("전처리 완료: %d 행 × %d 열", *df.shape)
    return df


# ──────────────────────────────────────────────
# 3. DB 적재 (upsert) — BUG FIX
# ──────────────────────────────────────────────
def upsert_to_db(df: pd.DataFrame) -> None:
    """
    [수정] SQLite upsert — 올바른 구현
      기존: df.to_sql(..., if_exists="append") 만 사용
            → 실행할 때마다 같은 데이터가 무한 중복 적재
      수정: 해당 PK 행을 먼저 DELETE 후 INSERT
            → 항상 최신 데이터 1건만 유지
    """
    engine = create_engine(DB_URL)
    # preprocess()에서 소문자+언더스코어 변환됨
    pk_col = PK_COLUMN.strip().lower().replace(" ", "_")

    with engine.begin() as conn:
        # 테이블 존재 여부 확인
        exists = conn.execute(
            text("SELECT name FROM sqlite_master WHERE type='table' AND name=:t"),
            {"t": TABLE_NAME},
        ).fetchone()

        if exists and pk_col in df.columns:
            # 수신된 PK 값에 해당하는 기존 행만 삭제
            pk_values = df[pk_col].dropna().tolist()
            if pk_values:
                # SQLite는 IN절 파라미터 바인딩이 번거로우므로 청크 처리
                chunk_size = 500
                deleted = 0
                for i in range(0, len(pk_values), chunk_size):
                    chunk = pk_values[i : i + chunk_size]
                    placeholders = ", ".join(f":v{j}" for j in range(len(chunk)))
                    params = {f"v{j}": v for j, v in enumerate(chunk)}
                    result = conn.execute(
                        text(f"DELETE FROM {TABLE_NAME} WHERE {pk_col} IN ({placeholders})"),
                        params,
                    )
                    deleted += result.rowcount
                log.info("기존 행 삭제: %d 건", deleted)

        # 신규 데이터 삽입
        df.to_sql(TABLE_NAME, conn, if_exists="append", index=False, method=None)

    log.info("DB 적재 완료: %d 행 → [%s]", len(df), TABLE_NAME)


# ──────────────────────────────────────────────
# 4. Excel 저장 (선택)
# ──────────────────────────────────────────────
def save_excel(df: pd.DataFrame) -> str:
    now      = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"공공서비스혜택정보_{now}.xlsx"

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="공공서비스혜택", index=False)
        ws = writer.sheets["공공서비스혜택"]

        from openpyxl.styles import Font, PatternFill, Alignment
        header_font  = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        header_fill  = PatternFill("solid", start_color="1F5C99")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        alt_fill     = PatternFill("solid", start_color="D6E4F0")

        for cell in ws[1]:
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = header_align

        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            for cell in row:
                if row_idx % 2 == 0:
                    cell.fill = alt_fill

        from openpyxl.utils import get_column_letter
        for col_idx, col_name in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = min(
                max(len(str(col_name)) * 2, 15), 50
            )

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    log.info("Excel 저장 완료: %s", filename)
    return filename


# ──────────────────────────────────────────────
# 5. 파이프라인 실행
# ──────────────────────────────────────────────
def run_pipeline() -> None:
    log.info("=" * 55)
    log.info("  파이프라인 시작: %s", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    log.info("=" * 55)

    raw = fetch_all()
    if not raw:
        log.error("데이터 수집 실패 — 파이프라인 중단")
        return

    df = preprocess(raw)
    upsert_to_db(df)

    if SAVE_EXCEL:
        save_excel(df)

    log.info("파이프라인 완료\n")


# ──────────────────────────────────────────────
# 6. 스케줄러
# ──────────────────────────────────────────────
def start_scheduler() -> None:
    scheduler = BlockingScheduler(timezone="Asia/Seoul")
    scheduler.add_job(
        run_pipeline,
        trigger="cron",
        hour=SCHEDULE_HR,
        minute=0,
        id="gov_benefits_job",
    )
    log.info("스케줄러 등록 완료 — 매일 %02d:00 (Asia/Seoul) 실행", SCHEDULE_HR)
    log.info("즉시 실행 후 대기 중... (Ctrl+C 로 종료)")
    run_pipeline()
    scheduler.start()


# ──────────────────────────────────────────────
if __name__ == "__main__":
    import sys
    if "--once" in sys.argv:
        run_pipeline()
    else:
        start_scheduler()
